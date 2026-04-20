import openpyxl
from docx import Document

from cover_letter.models import StudentProfile
from export.docx_exporter import export_letter
from export.xlsx_exporter import export_leads
from models import ContactInfo, LeadProfile, ScoreBreakdown


def _make_lead(name: str = "Techbolaget AB", tier: str = "STRONG") -> LeadProfile:
    return LeadProfile(
        company_name=name,
        website="https://example.com",
        city="Göteborg",
        tech_tags=["Java", "Spring"],
        contact=ContactInfo(email="jobb@example.com"),
        score=75.0,
        score_breakdown=ScoreBreakdown(tech_match=80, geo=70, contact=80, activity=70, stability=60, seniority=50),
        match_reason="Matchar Java, Göteborg",
        tier=tier,
    )


def _make_student() -> StudentProfile:
    return StudentProfile(
        name="Anna Lindqvist",
        education="YH Java-utvecklare",
        tech_stack=["Java"],
        experience="2 år support",
        languages=["svenska"],
        portfolio_url="https://github.com/anna",
        bio="Nyfiken",
    )


def test_xlsx_creates_file(tmp_path):
    path = tmp_path / "leads.xlsx"
    export_leads([_make_lead(), _make_lead("Bolaget AB", "GOOD")], path)
    assert path.exists()


def test_xlsx_has_correct_headers(tmp_path):
    path = tmp_path / "leads.xlsx"
    export_leads([_make_lead()], path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    assert "Företag" in headers
    assert "Score" in headers
    assert "Tier" in headers


def test_xlsx_has_data_row(tmp_path):
    path = tmp_path / "leads.xlsx"
    export_leads([_make_lead()], path)
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    assert ws.max_row == 2  # header + 1 data row
    assert ws.cell(2, 1).value == "Techbolaget AB"


def test_docx_creates_file(tmp_path):
    path = tmp_path / "brev.docx"
    export_letter("Hej! Jag söker LIA hos er.", _make_student(), "Techbolaget AB", path)
    assert path.exists()


def test_docx_contains_letter_text(tmp_path):
    path = tmp_path / "brev.docx"
    export_letter("Hej! Jag söker LIA hos er.", _make_student(), "Techbolaget AB", path)
    doc = Document(path)
    text = " ".join(p.text for p in doc.paragraphs)
    assert "Hej! Jag söker LIA hos er." in text
    assert "Techbolaget AB" in text
