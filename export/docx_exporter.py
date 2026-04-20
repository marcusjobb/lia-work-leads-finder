from datetime import date
from pathlib import Path

from docx import Document
from docx.shared import Pt

from cover_letter.models import StudentProfile


def export_letter(letter_text: str, student: StudentProfile, company_name: str, path: Path) -> None:
    doc = Document()

    header = doc.sections[0].header
    header_para = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
    header_para.text = f"{student.name}"
    if student.portfolio_url:
        header_para.text += f"  |  {student.portfolio_url}"
    header_para.runs[0].font.size = Pt(9)

    doc.add_paragraph(date.today().strftime("%d %B %Y"))
    doc.add_paragraph(f"Till: {company_name}")
    doc.add_paragraph("")

    for block in letter_text.strip().split("\n\n"):
        doc.add_paragraph(block.strip())

    doc.save(path)
